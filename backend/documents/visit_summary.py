"""Generates a DOCX visit summary for a patient after their appointment."""

from datetime import datetime
from pathlib import Path

from docx import Document

from backend.config import settings

OUTPUT_DIR = Path("./data/generated/visit_summaries")


def generate_visit_summary(
    patient_name: str,
    doctor_name: str,
    visit_date: str,
    doctor_notes: str,
    prescriptions: list[dict] | None = None,
) -> str:
    """Builds a DOCX visit summary and returns the saved file path."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.add_heading(settings.clinic_name, level=0)
    doc.add_paragraph(settings.clinic_address)
    doc.add_paragraph(f"Visit Summary — {visit_date}")

    doc.add_heading("Patient", level=1)
    doc.add_paragraph(patient_name)

    doc.add_heading("Attending Physician", level=1)
    doc.add_paragraph(doctor_name)

    doc.add_heading("Visit Notes", level=1)
    doc.add_paragraph(doctor_notes)

    if prescriptions:
        doc.add_heading("Prescriptions", level=1)
        for rx in prescriptions:
            doc.add_paragraph(
                f"{rx['medication_name']} — {rx['dosage']} — {rx['instructions']}",
                style="List Bullet",
            )

    filename = f"visit_summary_{patient_name.replace(' ', '_')}_{datetime.utcnow().strftime('%Y%m%d%H%M%S')}.docx"
    filepath = OUTPUT_DIR / filename
    doc.save(str(filepath))
    return str(filepath)
