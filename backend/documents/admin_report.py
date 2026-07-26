"""Generates a PPTX monthly operations report for the clinic admin."""

from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from backend.config import settings

OUTPUT_DIR = Path("./data/generated/admin_reports")


def generate_admin_report(metrics: dict) -> str:
    """Builds a PPTX report from the metrics dict returned by
    admin_reporting_agent.compile_monthly_metrics(), returns the saved path.
    """
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    prs = Presentation()

    # Title slide
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = f"{settings.clinic_name} — Monthly Report"
    title_slide.placeholders[1].text = datetime.utcnow().strftime("%B %Y")

    # Metrics slide
    metrics_slide = prs.slides.add_slide(prs.slide_layouts[1])
    metrics_slide.shapes.title.text = "Key Metrics"
    body = metrics_slide.placeholders[1].text_frame
    body.clear()

    rows = [
        f"Total patients: {metrics['total_patients']}",
        f"Appointments (last 30 days): {metrics['appointments_last_30d']}",
        f"No-shows: {metrics['no_shows']}",
        f"Claims filed: {metrics['claims_filed']}",
        f"Total claimed amount: ${metrics['total_claimed_amount']:.2f}",
    ]
    for i, row in enumerate(rows):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = row
        p.font.size = Pt(20)

    filename = f"admin_report_{datetime.utcnow().strftime('%Y%m%d%H%M%S')}.pptx"
    filepath = OUTPUT_DIR / filename
    prs.save(str(filepath))
    return str(filepath)
